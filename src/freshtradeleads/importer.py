from __future__ import annotations

import hashlib
import logging
from contextlib import nullcontext
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import delete, func, insert, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from .parsing import CLASSIFICATION_DESCRIPTIONS, clean, json_value, normalize_phone, parse_classifications, parse_date, row_digest
from .schema import bonds, classifications, contractor_classifications, contractors, import_rejections, license_snapshots, licenses, metadata, source_imports, workers_comp

LOG = logging.getLogger(__name__)
EXPECTED = {"LicenseNo", "BusinessName", "IssueDate", "Classifications(s)", "PrimaryStatus"}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _upsert(conn, table, rows: list[dict], keys: list[str], update_columns: list[str] | None = None):
    if not rows:
        return
    if conn.dialect.name == "postgresql":
        stmt = pg_insert(table).values(rows)
        updates = update_columns if update_columns is not None else [c.name for c in table.columns if c.name not in keys and not c.primary_key]
        if updates:
            stmt = stmt.on_conflict_do_update(index_elements=[table.c[k] for k in keys], set_={k: getattr(stmt.excluded, k) for k in updates})
        else:
            stmt = stmt.on_conflict_do_nothing(index_elements=[table.c[k] for k in keys])
        conn.execute(stmt)
    else:
        for row in rows:
            where = [table.c[k] == row[k] for k in keys]
            existing = conn.execute(select(table).where(*where)).mappings().first()
            if existing:
                updates = update_columns if update_columns is not None else [k for k in row if k not in keys]
                if updates:
                    conn.execute(update(table).where(*where).values(**{k: row[k] for k in updates if k in row}))
            else:
                conn.execute(insert(table).values(**row))


def _chunks(items: Iterable[tuple[int, dict[str, Any]]], size: int = 1000):
    batch = []
    for item in items:
        batch.append(item)
        if len(batch) >= size:
            yield batch
            batch = []
    if batch:
        yield batch


def import_xlsx(engine, file_path: str | Path, *, batch_size: int = 1000) -> dict[str, Any]:
    path = Path(file_path).resolve()
    digest = sha256_file(path)
    metadata.create_all(engine)
    with engine.begin() as conn:
        prior = conn.execute(select(source_imports.c.id).where(source_imports.c.sha256 == digest, source_imports.c.status == "complete").order_by(source_imports.c.id)).scalar()
        if prior:
            skipped = conn.execute(insert(source_imports).values(source_file=path.name, sha256=digest, file_size=path.stat().st_size, status="duplicate", duplicate_of_id=prior, completed_at=datetime.now(timezone.utc), imported_rows=0, rejected_rows=0).returning(source_imports.c.id)).scalar_one()
            return {"status": "duplicate", "import_id": skipped, "duplicate_of_id": prior, "sha256": digest}
        import_id = conn.execute(insert(source_imports).values(source_file=path.name, sha256=digest, file_size=path.stat().st_size, status="running").returning(source_imports.c.id)).scalar_one()

    total = imported = rejected = 0
    date_errors = Counter(); class_freq = Counter(); county_freq = Counter(); status_freq = Counter(); missing = Counter(); newest_issue = None
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
    if not EXPECTED.issubset(headers) or len(set(headers)) != len(headers):
        raise ValueError(f"Unexpected or duplicate workbook headers: {headers}")

    def records():
        nonlocal total
        for row_number, values in enumerate(iterator, 2):
            total += 1
            yield row_number, dict(zip(headers, values))

    try:
        connection = engine.connect()
        transaction = connection.begin()
        if connection.dialect.name == "postgresql":
            acquired = connection.execute(select(func.pg_try_advisory_xact_lock(1242026))).scalar()
            if not acquired:
                raise RuntimeError("another FreshTradeLeads import is already active")
        for batch in _chunks(records(), batch_size):
            good = []
            reject_rows = []
            for row_number, raw in batch:
                lic = clean(raw.get("LicenseNo"))
                business = clean(raw.get("BusinessName"))
                try:
                    issue = parse_date(raw.get("IssueDate"))
                    if not lic or not business or not issue:
                        raise ValueError("missing required license number, business name, or issue date")
                    parsed_dates = {h: parse_date(raw.get(h)) for h in ["LastUpdate","IssueDate","ReissueDate","ExpirationDate","InactivationDate","ReactivationDate","WCEffectiveDate","WCExpirationDate","WCCancellationDate","WCSuspendDate","CBEffectiveDate","CBCancellationDate","WBEffectiveDate","WBCancellationDate","DBEffectiveDate","DBCancellationDate","DateRequired"]}
                    codes = parse_classifications(raw.get("Classifications(s)"))
                    if not codes:
                        raise ValueError("missing classification")
                    good.append((row_number, raw, lic, business, issue, parsed_dates, codes, row_digest(raw)))
                except Exception as exc:
                    rejected += 1
                    reject_rows.append({"import_id": import_id, "row_number": row_number, "license_number_raw": lic, "reason": str(exc), "raw_record": {k: json_value(v) for k,v in raw.items()}})
            with nullcontext(connection) as conn:
                if reject_rows: conn.execute(insert(import_rejections), reject_rows)
                contractor_rows = [{"license_number": x[2], "business_name": x[3], "business_name_2": clean(x[1].get("BUS-NAME-2")), "full_business_name": clean(x[1].get("FullBusinessName")), "business_type": clean(x[1].get("BusinessType")), "mailing_address": clean(x[1].get("MailingAddress")), "city": clean(x[1].get("City")), "state": clean(x[1].get("State")), "county": clean(x[1].get("County")), "zip_code": clean(x[1].get("ZIPCode")), "country": clean(x[1].get("country")), "phone_raw": clean(x[1].get("BusinessPhone")), "phone_normalized": normalize_phone(x[1].get("BusinessPhone")), "last_seen_at": datetime.now(timezone.utc)} for x in good]
                _upsert(conn, contractors, contractor_rows, ["license_number"], ["business_name","business_name_2","full_business_name","business_type","mailing_address","city","state","county","zip_code","country","phone_raw","phone_normalized","last_seen_at"])
                ids = dict(conn.execute(select(contractors.c.license_number, contractors.c.id).where(contractors.c.license_number.in_([x[2] for x in good]))).all())
                license_rows = [{"license_number": x[2], "contractor_id": ids[x[2]], "last_source_update": x[5]["LastUpdate"], "original_issue_date": x[4], "reissue_date": x[5]["ReissueDate"], "expiration_date": x[5]["ExpirationDate"], "inactivation_date": x[5]["InactivationDate"], "reactivation_date": x[5]["ReactivationDate"], "primary_status": clean(x[1].get("PrimaryStatus")), "secondary_status": clean(x[1].get("SecondaryStatus")), "pending_suspension": clean(x[1].get("PendingSuspension")), "pending_class_removal": clean(x[1].get("PendingClassRemoval")), "pending_class_replace": clean(x[1].get("PendingClassReplace")), "asbestos_registration": clean(x[1].get("AsbestosReg")), "classifications_raw": clean(x[1].get("Classifications(s)")), "first_seen_import_id": import_id, "last_seen_import_id": import_id, "current_row_hash": x[7]} for x in good]
                _upsert(conn, licenses, license_rows, ["license_number"], [c for c in license_rows[0] if c not in {"license_number","contractor_id","first_seen_import_id"}] if license_rows else [])
                all_codes = sorted({c for x in good for c in x[6]})
                _upsert(conn, classifications, [{"code": c, "description": CLASSIFICATION_DESCRIPTIONS.get(c)} for c in all_codes], ["code"], ["description"])
                lic_nums = [x[2] for x in good]
                if lic_nums:
                    conn.execute(delete(contractor_classifications).where(contractor_classifications.c.license_number.in_(lic_nums)))
                    conn.execute(insert(contractor_classifications), [{"license_number": x[2], "classification_code": code, "position": pos} for x in good for pos,code in enumerate(x[6],1)])
                wc_rows = [{"license_number": x[2], "coverage_type": clean(x[1].get("WorkersCompCoverageType")), "carrier": clean(x[1].get("WCInsuranceCompany")), "policy_number": clean(x[1].get("WCPolicyNumber")), "effective_date": x[5]["WCEffectiveDate"], "expiration_date": x[5]["WCExpirationDate"], "cancellation_date": x[5]["WCCancellationDate"], "suspension_date": x[5]["WCSuspendDate"]} for x in good]
                _upsert(conn, workers_comp, wc_rows, ["license_number"])
                bond_rows=[]
                for x in good:
                    r=x[1]; d=x[5]
                    for kind,prefix in [("contractor","CB"),("worker","WB"),("disciplinary","DB")]:
                        if clean(r.get(prefix+"SuretyCompany")) or clean(r.get(prefix+"Number")) or r.get(prefix+"Amount") is not None:
                            bond_rows.append({"license_number":x[2],"bond_type":kind,"surety_company":clean(r.get(prefix+"SuretyCompany")),"bond_number":clean(r.get(prefix+"Number")),"effective_date":d.get(prefix+"EffectiveDate"),"cancellation_date":d.get(prefix+"CancellationDate"),"amount":r.get(prefix+"Amount"),"date_required":d.get("DateRequired") if kind=="disciplinary" else None,"case_region":clean(r.get("DiscpCaseRegion")) if kind=="disciplinary" else None,"reason":clean(r.get("DBBondReason")) if kind=="disciplinary" else None,"case_number":clean(r.get("DBCaseNo")) if kind=="disciplinary" else None})
                _upsert(conn, bonds, bond_rows, ["license_number","bond_type"])
                conn.execute(insert(license_snapshots), [{"import_id":import_id,"license_number":x[2],"row_number":x[0],"row_hash":x[7],"raw_record":{k:json_value(v) for k,v in x[1].items()}} for x in good])
            imported += len(good)
            for x in good:
                newest_issue = max(newest_issue, x[4]) if newest_issue else x[4]
                class_freq.update(x[6]); county_freq[clean(x[1].get("County")) or "(missing)"] += 1; status_freq[clean(x[1].get("PrimaryStatus")) or "(missing)"] += 1
                for key in ["BusinessName","BusinessPhone","IssueDate"]:
                    if not clean(x[1].get(key)): missing[key] += 1
            if imported % 10000 < batch_size: LOG.info("Imported %s of %s rows", imported, ws.max_row - 1)
        stats={"classification_frequency":class_freq.most_common(),"county_frequency":county_freq.most_common(),"license_status_frequency":status_freq.most_common(),"missing_fields":dict(missing),"invalid_dates":dict(date_errors)}
        connection.execute(update(source_imports).where(source_imports.c.id==import_id).values(status="complete",completed_at=datetime.now(timezone.utc),sheet_name=ws.title,workbook_rows=total,imported_rows=imported,rejected_rows=rejected,newest_issue_date=newest_issue,statistics=stats))
        transaction.commit(); connection.close()
        return {"status":"complete","import_id":import_id,"sha256":digest,"sheet":ws.title,"workbook_rows":total,"imported_rows":imported,"rejected_rows":rejected,"newest_issue_date":str(newest_issue)}
    except Exception as exc:
        if "transaction" in locals() and transaction.is_active: transaction.rollback()
        if "connection" in locals(): connection.close()
        with engine.begin() as conn: conn.execute(update(source_imports).where(source_imports.c.id==import_id).values(status="failed",completed_at=datetime.now(timezone.utc),workbook_rows=total,imported_rows=imported,rejected_rows=rejected,error_message=str(exc)))
        raise
