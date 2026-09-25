from __future__ import annotations

import fcntl
import hashlib
import html
import http.cookiejar
import logging
import os
import re
import tempfile
import urllib.request
import urllib.parse
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import insert, select, update

from .importer import import_xlsx
from .schema import metadata, source_fetches, source_imports

LOG = logging.getLogger(__name__)
PORTAL_URL = "https://www.cslb.ca.gov/onlineservices/dataportal/ContractorList"
DOWNLOAD_URL = "https://www.cslb.ca.gov/OnlineServices/DataPortal/DownLoadFile.ashx?fName=MasterLicenseData&type=X"
CRITICAL_COLUMNS = {"LicenseNo", "BusinessName", "IssueDate", "PrimaryStatus", "Classifications(s)"}
MIN_PRODUCTION_BYTES = 1_000_000
MIN_PRODUCTION_ROWS = 100_000


class SourceValidationError(ValueError):
    pass


class FetchAlreadyRunning(RuntimeError):
    pass


@dataclass
class WorkbookInfo:
    sheet_name: str
    data_rows: int
    columns: list[str]


@contextmanager
def acquisition_lock(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a+") as handle:
        try:
            fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise FetchAlreadyRunning(f"another source acquisition is active ({path})") from exc
        yield


def discover_source_date(portal_url: str = PORTAL_URL, *, timeout: int = 60) -> date:
    opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    request = urllib.request.Request(portal_url, headers={"User-Agent": "FreshTradeLeads/0.2 (+CSLB public data importer)"})
    with opener.open(request, timeout=timeout) as response:
        body = response.read().decode("utf-8", errors="replace")
    hidden={name:html.unescape(value) for name,value in re.findall(r'<input[^>]+name="([^"]+)"[^>]+value="([^"]*)"',body) if name.startswith("__")}
    hidden["ctl00$MainContent$ddlStatus"]="M"
    post=urllib.request.Request(portal_url,urllib.parse.urlencode(hidden).encode(),headers={"Content-Type":"application/x-www-form-urlencoded","User-Agent":"FreshTradeLeads/0.2 (+CSLB public data importer)"})
    with opener.open(post,timeout=timeout) as response:
        body=response.read().decode("utf-8",errors="replace")
    match = re.search(r"Updated\s+as\s+of\s+(\d{1,2}/\d{1,2}/\d{4})", body, re.I)
    if not match:
        raise SourceValidationError("official portal did not expose an 'Updated as of' source date")
    return datetime.strptime(match.group(1), "%m/%d/%Y").date()


def validate_xlsx(path: Path, *, min_bytes: int = MIN_PRODUCTION_BYTES, min_rows: int = MIN_PRODUCTION_ROWS) -> WorkbookInfo:
    if not path.exists() or path.stat().st_size < min_bytes:
        raise SourceValidationError(f"download is implausibly small: {path.stat().st_size if path.exists() else 0} bytes")
    if not zipfile.is_zipfile(path):
        prefix = path.read_bytes()[:80]
        raise SourceValidationError(f"download is not a ZIP/XLSX container; prefix={prefix!r}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if not workbook.sheetnames:
            raise SourceValidationError("workbook has no worksheets")
        if "CSLBMasterLicenseData" not in workbook.sheetnames:
            raise SourceValidationError(f"expected sheet CSLBMasterLicenseData not found; sheets={workbook.sheetnames}")
        worksheet = workbook["CSLBMasterLicenseData"]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = [str(value).strip() if value is not None else "" for value in header]
        missing = sorted(CRITICAL_COLUMNS - set(columns))
        if missing:
            raise SourceValidationError(f"workbook is missing critical columns: {', '.join(missing)}")
        rows = worksheet.max_row - 1
        if rows < min_rows:
            raise SourceValidationError(f"workbook has only {rows:,} data rows; expected at least {min_rows:,}")
        return WorkbookInfo(worksheet.title, rows, columns)
    except SourceValidationError:
        raise
    except Exception as exc:
        raise SourceValidationError(f"XLSX could not be opened: {exc}") from exc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _download(url: str, target: Path, *, timeout: int = 300) -> tuple[dict, int]:
    request = urllib.request.Request(url, headers={"User-Agent": "FreshTradeLeads/0.2 (+CSLB public data importer)", "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/vnd.ms-excel"})
    with urllib.request.urlopen(request, timeout=timeout) as response, target.open("wb") as output:
        headers = {key.lower(): value for key, value in response.headers.items()}
        total = 0
        while block := response.read(1024 * 1024):
            output.write(block); total += len(block)
        output.flush(); os.fsync(output.fileno())
        return {"status": response.status, "etag": headers.get("etag"), "last_modified": headers.get("last-modified"), "content_length": int(headers["content-length"]) if headers.get("content-length", "").isdigit() else None}, total


def fetch_source(engine, raw_dir: Path, *, portal_url: str = PORTAL_URL, download_url: str = DOWNLOAD_URL, min_bytes: int = MIN_PRODUCTION_BYTES, min_rows: int = MIN_PRODUCTION_ROWS, run_pipeline=None) -> dict:
    metadata.create_all(engine)
    raw_dir.mkdir(parents=True, exist_ok=True)
    lock_path = raw_dir.parent / ".source-fetch.lock"
    with acquisition_lock(lock_path):
        with engine.begin() as connection:
            fetch_id = connection.execute(insert(source_fetches).values(portal_url=portal_url, download_url=download_url, status="running").returning(source_fetches.c.id)).scalar_one()
        temp_path = None
        try:
            source_date = discover_source_date(portal_url)
            descriptor, temp_name = tempfile.mkstemp(prefix=".MasterLicenseData.", suffix=".part.xlsx", dir=raw_dir)
            os.close(descriptor); temp_path = Path(temp_name)
            http, downloaded = _download(download_url, temp_path)
            info = validate_xlsx(temp_path, min_bytes=min_bytes, min_rows=min_rows)
            digest = _sha256(temp_path)
            with engine.connect() as connection:
                prior = connection.execute(select(source_imports.c.id, source_imports.c.source_file).where(source_imports.c.sha256 == digest, source_imports.c.status == "complete").order_by(source_imports.c.id.desc()).limit(1)).mappings().first()
            common = dict(completed_at=datetime.now(timezone.utc), portal_source_date=source_date, http_status=http["status"], etag=http["etag"], last_modified=http["last_modified"], content_length=http["content_length"], downloaded_bytes=downloaded, sha256=digest)
            if prior:
                temp_path.unlink(); temp_path = None
                with engine.begin() as connection: connection.execute(update(source_fetches).where(source_fetches.c.id == fetch_id).values(status="unchanged", message=f"same hash as import {prior['id']} ({prior['source_file']})", **common))
                return {"status":"unchanged","fetch_id":fetch_id,"source_date":source_date.isoformat(),"sha256":digest,"downloaded_bytes":downloaded,"matching_import_id":prior["id"],"workbook_rows":info.data_rows,"http":http}
            destination = raw_dir / f"{source_date.isoformat()}_MasterLicenseData.xlsx"
            if destination.exists():
                destination = raw_dir / f"{source_date.isoformat()}_MasterLicenseData_{digest[:8]}.xlsx"
            os.replace(temp_path, destination); temp_path = None
            import_result = import_xlsx(engine, destination)
            pipeline_result = run_pipeline(import_result["import_id"], source_date) if run_pipeline else None
            with engine.begin() as connection: connection.execute(update(source_fetches).where(source_fetches.c.id == fetch_id).values(status="imported", saved_path=str(destination), import_id=import_result["import_id"], message="download validated and imported", **common))
            return {"status":"imported","fetch_id":fetch_id,"source_date":source_date.isoformat(),"sha256":digest,"downloaded_bytes":downloaded,"saved_path":str(destination),"workbook_rows":info.data_rows,"import":import_result,"pipeline":pipeline_result,"http":http}
        except Exception as exc:
            if temp_path and temp_path.exists(): temp_path.unlink()
            with engine.begin() as connection: connection.execute(update(source_fetches).where(source_fetches.c.id == fetch_id).values(status="failed", completed_at=datetime.now(timezone.utc), message=str(exc)))
            LOG.exception("CSLB source acquisition failed")
            raise
