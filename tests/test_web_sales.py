from __future__ import annotations

import json
import hashlib
import hmac
import time
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from freshtradeleads.config import WebSettings
from freshtradeleads.importer import import_xlsx
from freshtradeleads.orders import authorize_download, create_order, fulfill_verified_event, order_status
from freshtradeleads.products import REPORT_COLUMNS, current_preview, historical_sample, inventory
from freshtradeleads.schema import download_tokens, funnel_events, order_leads, orders
from freshtradeleads.stripe_gateway import CheckoutResult
from freshtradeleads.stripe_gateway import StripeGateway
from freshtradeleads.web import create_app
from conftest import HEADERS


class MockStripe:
    def __init__(self): self.calls=[]
    def create_checkout(self, **kwargs):
        self.calls.append(kwargs); return CheckoutResult(f"cs_test_{len(self.calls)}",f"https://checkout.stripe.test/{len(self.calls)}")
    def verify_webhook(self,payload,signature):
        assert signature == "valid"; return json.loads(payload)


@pytest.fixture
def sales_data(engine,tmp_path):
    path=tmp_path/"sales.xlsx"; wb=Workbook(); ws=wb.active; ws.title="CSLBMasterLicenseData"; ws.append(HEADERS)
    anchor=date(2026,9,2)
    def add(i,issued):
        d={h:None for h in HEADERS}; d.update({"LicenseNo":str(200000+i),"LastUpdate":"09/03/2026","BusinessName":f"COMPLETE ELECTRIC {i}","FullBusinessName":f"COMPLETE ELECTRIC {i}","MailingAddress":f"{i} MAIN ST","City":"LOS ANGELES","State":"CA","County":"Los Angeles","ZIPCode":"90001","BusinessPhone":"(213) 555-0100","BusinessType":"Corporation","IssueDate":issued.strftime("%m/%d/%Y"),"ExpirationDate":"09/30/2028","PrimaryStatus":"CLEAR","Classifications(s)":"C10","WorkersCompCoverageType":"Insurance","WCInsuranceCompany":"TEST CARRIER","CBSuretyCompany":"TEST SURETY","CBAmount":25000})
        ws.append([d[h] for h in HEADERS])
    for i in range(30): add(i,anchor-timedelta(days=i))
    for i,days in enumerate((90,95,100,110,120),30): add(i,anchor-timedelta(days=days))
    wb.save(path); import_xlsx(engine,path); return anchor


@pytest.fixture
def settings(tmp_path,engine):
    return WebSettings(str(engine.url),tmp_path,"http://testserver","test-session-secret",None,None,7,None,False)


def test_paid_skus_preview_and_sample(engine,sales_data):
    with engine.begin() as conn:
        snapshot,items=inventory(conn); counts={x["sku"]:x["lead_count"] for x in items}
        preview=current_preview(conn); _,sample=historical_sample(conn)
    assert counts == {"starter":10,"fresh-pack":14,"expanded":30}
    assert all(set(row) == {"lead_number","trade","region","issued","status","phone_included"} for row in preview)
    assert all(forbidden not in str(preview).lower() for forbidden in ("business name","license number","address","zip","city"))
    assert len(sample)==5 and all(90 <= (sales_data-r["Original Issue Date"]).days <= 120 for r in sample)
    assert all(r["Business Name"] and r["Phone"] and r["Address"] and r["License Number"] for r in sample)


@pytest.mark.parametrize("sku,amount,count",[("starter",1500,10),("fresh-pack",2900,14),("expanded",5900,30)])
def test_order_pinning_checkout_and_fulfillment(engine,sales_data,tmp_path,sku,amount,count):
    gateway=MockStripe(); oid,token,checkout=create_order(engine,sku,utm={"utm_source":"test"},visitor_id="visitor",base_url="http://testserver",gateway=gateway)
    order=order_status(engine,oid); assert order["price_cents"]==amount and order["pinned_lead_count"]==count and order["status"]=="checkout_pending"
    assert gateway.calls[0]["amount"]==amount and gateway.calls[0]["sku"]==sku
    with engine.connect() as conn: pinned=list(conn.execute(select(order_leads.c.license_number).where(order_leads.c.order_id==oid).order_by(order_leads.c.position)).scalars())
    event={"id":f"evt_{sku}","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"paid","amount_total":amount,"currency":"usd","payment_intent":"pi_test","metadata":{"order_id":oid,"sku":sku}}}}
    result=fulfill_verified_event(engine,event,data_dir=tmp_path); assert result["status"]=="fulfilled"
    assert fulfill_verified_event(engine,event,data_dir=tmp_path)["duplicate"] is True
    xlsx=authorize_download(engine,token,"xlsx",data_dir=tmp_path); csv=authorize_download(engine,token,"csv",data_dir=tmp_path)
    assert xlsx.exists() and csv.exists()
    assert load_workbook(xlsx,read_only=True)["Leads"].max_row-1 == count
    assert sum(1 for line in csv.read_text(encoding="utf-8-sig").splitlines() if line and not line.startswith("#"))-1 == count
    with engine.begin() as conn: conn.execute(update(download_tokens).where(download_tokens.c.order_id==oid).values(expires_at=datetime.now(timezone.utc)-timedelta(seconds=1)))
    assert authorize_download(engine,token,"xlsx",data_dir=tmp_path) is None
    with engine.connect() as conn: assert list(conn.execute(select(order_leads.c.license_number).where(order_leads.c.order_id==oid).order_by(order_leads.c.position)).scalars()) == pinned


def test_success_url_cannot_fulfill_and_utm_persists(engine,sales_data,settings):
    gateway=MockStripe(); app=create_app(engine=engine,settings=settings,gateway=gateway); client=TestClient(app)
    page=client.get("/?utm_source=facebook&utm_campaign=c10"); assert page.status_code==200
    csrf=client.cookies.get("session"); assert csrf
    from itsdangerous import TimestampSigner
    # Submit through the rendered form to preserve the signed session and CSRF value.
    import re
    token=re.search(r'name="csrf" value="([^"]+)"',page.text).group(1)
    response=client.post("/checkout/fresh-pack",data={"csrf":token},follow_redirects=False); assert response.status_code==303
    with engine.connect() as conn:
        order=conn.execute(select(orders).order_by(orders.c.created_at.desc())).mappings().first()
        assert order["utm_attribution"]["utm_source"]=="facebook"
    returned=client.get(f"/order/{order['id']}/success?token=invalid&session_id={order['stripe_checkout_session_id']}")
    assert returned.status_code==200 and order_status(engine,order["id"])["status"]=="checkout_pending"
    assert client.get("/download/invalid/xlsx").status_code==404


def test_real_stripe_signature_validation():
    payload=b'{"id":"evt_sig","type":"test.event","data":{"object":{}}}'
    secret="whsec_test_fixture"; timestamp=int(time.time())
    digest=hmac.new(secret.encode(),f"{timestamp}.".encode()+payload,hashlib.sha256).hexdigest()
    event=StripeGateway("sk_test_fixture",secret,{"starter":"price_test"}).verify_webhook(payload,f"t={timestamp},v1={digest}")
    assert event["id"]=="evt_sig"
    with pytest.raises(Exception): StripeGateway("sk_test_fixture",secret,{"starter":"price_test"}).verify_webhook(payload,f"t={timestamp},v1=bad")


def test_checkout_disables_managed_payments_and_uses_card(monkeypatch):
    captured={}
    class Session:
        id="cs_test"; url="https://checkout.stripe.test/session"
    def fake_create(**kwargs): captured.update(kwargs); return Session()
    monkeypatch.setattr("stripe.checkout.Session.create",fake_create)
    gateway=StripeGateway("sk_test_fixture","whsec_fixture",{"starter":"price_starter"})
    gateway.create_checkout(order_id="order-123",sku="starter",name="Starter",description="Ten leads",amount=1500,currency="usd",success_url="http://test/success",cancel_url="http://test/cancel")
    assert captured["line_items"]==[{"price":"price_starter","quantity":1}]
    assert captured["managed_payments"]=={"enabled":False}
    assert captured["payment_method_types"]==["card"]
    assert captured["idempotency_key"]=="order-123"
    assert "price_data" not in captured["line_items"][0]
    assert captured["allow_promotion_codes"] is True


def test_allowlisted_auto_promotion_and_discount_persistence(engine,sales_data,tmp_path):
    gateway=MockStripe()
    oid,_,checkout=create_order(engine,"fresh-pack",utm={},visitor_id=None,base_url="https://example.test",gateway=gateway,promotion_code="FRESH20",promotion_id="promo_approved")
    order=order_status(engine,oid)
    assert order["price_cents"]==2900 and order["final_price_cents"]==2320 and order["discount_cents"]==580
    assert gateway.calls[0]["promotion_id"]=="promo_approved" and gateway.calls[0]["allow_promotion_codes"] is False
    event={"id":"evt_discount","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"paid","amount_total":2320,"total_details":{"amount_discount":580},"discounts":[{"promotion_code":"promo_approved"}],"currency":"usd","metadata":{"order_id":oid}}}}
    assert fulfill_verified_event(engine,event,data_dir=tmp_path,approved_promotion_id="promo_approved",approved_promotion_code="FRESH20")["status"]=="fulfilled"
    assert order_status(engine,oid)["final_price_cents"]==2320


def test_mismatched_or_arbitrary_discount_is_rejected(engine,sales_data,tmp_path):
    gateway=MockStripe(); oid,_,checkout=create_order(engine,"starter",utm={},visitor_id=None,base_url="https://example.test",gateway=gateway)
    event={"id":"evt_bad_discount","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"paid","amount_total":1300,"total_details":{"amount_discount":200},"currency":"usd","metadata":{"order_id":oid}}}}
    with pytest.raises(ValueError,match="unapproved discount"): fulfill_verified_event(engine,event,data_dir=tmp_path)
    event["id"]="evt_wrong_code"; event["data"]["object"].update({"amount_total":1200,"total_details":{"amount_discount":300},"discounts":[{"promotion_code":"promo_other"}]})
    with pytest.raises(ValueError,match="promotion code"): fulfill_verified_event(engine,event,data_dir=tmp_path,approved_promotion_id="promo_approved")


def test_public_cookie_host_and_promo_allowlist(engine,sales_data,settings):
    public=replace(settings,base_url="https://leads.example.com",secure_cookies=True,allowed_hosts=("leads.example.com","testserver"),promotion_code="FRESH20",stripe_promotion_id="promo_approved")
    gateway=MockStripe(); client=TestClient(create_app(engine=engine,settings=public,gateway=gateway),base_url="https://leads.example.com")
    page=client.get("/?promo=NOTALLOWED"); assert page.status_code==200
    import re
    csrf=re.search(r'name="csrf" value="([^"]+)"',page.text).group(1)
    client.post("/checkout/starter",data={"csrf":csrf},follow_redirects=False)
    assert gateway.calls[-1]["promotion_id"] is None
    page=client.get("/?promo=FRESH20")
    assert "Launch offer" in page.text and "secure" in page.headers["set-cookie"].lower()
    csrf=re.search(r'name="csrf" value="([^"]+)"',page.text).group(1)
    client.post("/checkout/starter",data={"csrf":csrf},follow_redirects=False)
    assert gateway.calls[-1]["promotion_id"]=="promo_approved"
    assert client.get("/",headers={"host":"evil.example"}).status_code==400


def test_expired_and_unpaid_events_do_not_fulfill(engine,sales_data,tmp_path):
    gateway=MockStripe(); oid,_,checkout=create_order(engine,"starter",utm={},visitor_id=None,base_url="http://test",gateway=gateway)
    unpaid={"id":"evt_unpaid","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"unpaid","amount_total":1500,"currency":"usd","metadata":{"order_id":oid}}}}
    assert fulfill_verified_event(engine,unpaid,data_dir=tmp_path)["status"]=="payment_pending"
    assert order_status(engine,oid)["status"]=="checkout_pending"
    expired={"id":"evt_expired","type":"checkout.session.expired","data":{"object":{"id":checkout.id,"metadata":{"order_id":oid}}}}
    assert fulfill_verified_event(engine,expired,data_dir=tmp_path)["status"]=="processed"
    assert order_status(engine,oid)["status"]=="cancelled"


def test_partial_and_full_refunds_are_idempotent_and_revoke_download(engine,sales_data,tmp_path):
    gateway=MockStripe()
    oid,token,checkout=create_order(engine,"starter",utm={"utm_campaign":"refund-test"},visitor_id=None,base_url="http://test",gateway=gateway)
    paid={"id":"evt_paid_refund","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"paid","amount_total":1500,"currency":"usd","payment_intent":"pi_refund_test","metadata":{"order_id":oid}}}}
    assert fulfill_verified_event(engine,paid,data_dir=tmp_path)["status"]=="fulfilled"
    assert authorize_download(engine,token,"xlsx",data_dir=tmp_path)

    partial={"id":"evt_partial_refund","type":"charge.refunded","data":{"object":{"payment_intent":"pi_refund_test","amount":1500,"amount_refunded":500}}}
    assert fulfill_verified_event(engine,partial,data_dir=tmp_path)["status"]=="partially_refunded"
    assert fulfill_verified_event(engine,partial,data_dir=tmp_path)["duplicate"] is True
    order=order_status(engine,oid)
    assert order["status"]=="partially_refunded" and order["refund_amount_cents"]==500
    assert authorize_download(engine,token,"csv",data_dir=tmp_path)

    full={"id":"evt_full_refund","type":"charge.refunded","data":{"object":{"payment_intent":"pi_refund_test","amount":1500,"amount_refunded":1500}}}
    assert fulfill_verified_event(engine,full,data_dir=tmp_path)["status"]=="refunded"
    assert fulfill_verified_event(engine,full,data_dir=tmp_path)["duplicate"] is True
    order=order_status(engine,oid)
    assert order["status"]=="refunded" and order["refund_amount_cents"]==1500 and order["refunded_at"]
    assert authorize_download(engine,token,"xlsx",data_dir=tmp_path) is None


def test_live_stripe_key_is_accepted_for_production_settings(tmp_path):
    settings=WebSettings(
        "sqlite://",tmp_path,"https://leads.example.com","session-secret",
        "sk_live_fixture","whsec_fixture",stripe_prices={
            "starter":"price_starter","fresh-pack":"price_fresh","expanded":"price_expanded"
        },
    )
    assert settings.stripe_ready


def test_meta_pixel_events_and_privacy_disclosure(engine,sales_data,settings,tmp_path):
    tracked=replace(settings,meta_pixel_id="123456789012345")
    gateway=MockStripe()
    client=TestClient(create_app(engine=engine,settings=tracked,gateway=gateway))
    landing=client.get("/")
    assert 'data-meta-pixel-id="123456789012345"' in landing.text
    assert '/static/meta-pixel.js' in landing.text
    assert "data-meta-checkout" in landing.text
    csp=landing.headers["content-security-policy"]
    assert "https://connect.facebook.net" in csp and "https://www.facebook.com" in csp
    script=client.get("/static/meta-pixel.js").text
    assert '"PageView"' in script and '"InitiateCheckout"' in script and '"Purchase"' in script
    privacy=client.get("/privacy").text
    assert "Meta Pixel" in privacy and "authoritative source" in privacy

    oid,token,checkout=create_order(engine,"starter",utm={"utm_source":"facebook"},visitor_id="pixel-test",base_url="http://testserver",gateway=gateway)
    paid={"id":"evt_meta_purchase","type":"checkout.session.completed","data":{"object":{"id":checkout.id,"payment_status":"paid","amount_total":1500,"currency":"usd","payment_intent":"pi_meta","metadata":{"order_id":oid}}}}
    fulfill_verified_event(engine,paid,data_dir=tmp_path)
    success=client.get(f"/order/{oid}/success?token={token}&session_id={checkout.id}")
    assert "data-meta-purchase" in success.text
    assert f'data-meta-order="{oid}"' in success.text
    assert 'data-meta-value="15.00"' in success.text


def test_meta_pixel_is_absent_when_not_configured(engine,sales_data,settings):
    client=TestClient(create_app(engine=engine,settings=settings,gateway=MockStripe()))
    page=client.get("/")
    assert "data-meta-pixel-id" not in page.text and "/static/meta-pixel.js" not in page.text
