from openpyxl import load_workbook
from sqlalchemy import func, select
from freshtradeleads.importer import import_xlsx
from freshtradeleads.schema import import_rejections, licenses


def test_duplicate_license_in_distinct_snapshot_updates_instead_of_duplicates(engine,workbook):
    import_xlsx(engine,workbook)
    wb=load_workbook(workbook); ws=wb.active; ws["C2"]="Alpha Electric Updated"; changed=workbook.parent/"changed.xlsx"; wb.save(changed)
    import_xlsx(engine,changed)
    with engine.connect() as c:
        assert c.scalar(select(func.count()).select_from(licenses)) == 3
        assert c.scalar(select(func.count()).select_from(import_rejections)) == 0

