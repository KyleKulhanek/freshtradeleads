from datetime import date
from openpyxl import load_workbook
from sqlalchemy import func, select

from freshtradeleads.exporter import export_xlsx
from freshtradeleads.importer import import_xlsx, sha256_file
from freshtradeleads.query import lead_query
from freshtradeleads.schema import contractor_classifications, contractors, licenses, source_imports


def test_import_is_idempotent_and_hash_is_recorded(engine,workbook):
    first=import_xlsx(engine,workbook,batch_size=2); second=import_xlsx(engine,workbook,batch_size=2)
    assert first["imported_rows"] == 3 and first["rejected_rows"] == 0
    assert second["status"] == "duplicate" and second["duplicate_of_id"] == first["import_id"]
    with engine.connect() as c:
        assert c.scalar(select(func.count()).select_from(contractors)) == 3
        assert c.scalar(select(func.count()).select_from(licenses)) == 3
        assert c.scalar(select(func.count()).select_from(source_imports)) == 2
        assert c.scalar(select(source_imports.c.sha256).where(source_imports.c.id==first["import_id"])) == sha256_file(workbook)


def test_core_filtering(engine,workbook):
    import_xlsx(engine,workbook)
    with engine.connect() as c:
        rows=c.execute(lead_query(classification="C10",county="Los Angeles",days=30,anchor_date=date(2026,9,2),dialect_name="sqlite")).mappings().all()
    assert [r["License Number"] for r in rows] == ["001001"]


def test_days_window_contains_exactly_n_calendar_dates(engine,workbook):
    import_xlsx(engine,workbook)
    with engine.connect() as c:
        rows=c.execute(lead_query(days=4,anchor_date=date(2026,9,2),dialect_name="sqlite")).mappings().all()
    assert {r["License Number"] for r in rows} == {"001001","001002"}


def test_xlsx_export(engine,workbook,tmp_path):
    result=import_xlsx(engine,workbook)
    with engine.connect() as c: rows=[dict(r) for r in c.execute(lead_query(classification="C10",days=90,anchor_date=date(2026,9,2),dialect_name="sqlite")).mappings()]
    out=tmp_path/"report.xlsx"; export_xlsx(rows,list(rows[0]),out,{"generated_at":"2026-09-03T00:00:00Z","source_as_of":"2026-09-02","source_sha256":result["sha256"]})
    wb=load_workbook(out,read_only=False); ws=wb["Leads"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws.max_row == 3
