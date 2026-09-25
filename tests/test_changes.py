from openpyxl import load_workbook

from freshtradeleads.changes import compare_snapshots
from freshtradeleads.importer import import_xlsx
from conftest import HEADERS


def test_snapshot_change_categories(engine,workbook,tmp_path):
    first=import_xlsx(engine,workbook)
    wb=load_workbook(workbook); ws=wb.active
    ws["V2"]="Contr Bond Susp"       # status
    ws["X2"]="B| C10| C36"          # classification
    ws["Z2"]="Workers' Compensation Insurance"; ws["AA2"]="NEW CARRIER"  # WC
    ws["AG2"]="NEW SURETY"; ws["AK2"]=30000                              # bond
    base={h:None for h in HEADERS}; base.update({"LicenseNo":"001004","LastUpdate":"09/04/2026","BusinessName":"Delta HVAC","MailingAddress":"2 MAIN ST","City":"LOS ANGELES","State":"CA","County":"Los Angeles","ZIPCode":"90001","BusinessPhone":"2135550104","BusinessType":"Corporation","IssueDate":"09/03/2026","ExpirationDate":"09/30/2028","PrimaryStatus":"CLEAR","Classifications(s)":"C20","WorkersCompCoverageType":"Exempt","WCEffectiveDate":"09/03/2026","CBSuretyCompany":"TEST SURETY","CBNumber":"B4","CBEffectiveDate":"09/03/2026","CBAmount":25000})
    ws.append([base[h] for h in HEADERS]); changed=workbook.parent/"snapshot2.xlsx"; wb.save(changed)
    second=import_xlsx(engine,changed)
    report=compare_snapshots(engine,second["import_id"],tmp_path)
    for kind in ["new-source-record","newly-licensed","status-change","classification-change","workers-comp-change","bond-change"]:
        assert report["counts"][kind]>=1
    assert report["counts"]["source-record-removed"]==0
    assert (tmp_path/f"change_report_2026-09-03_import_{second['import_id']}.json").exists()


def test_first_snapshot_is_baseline(engine,workbook,tmp_path):
    first=import_xlsx(engine,workbook)
    report=compare_snapshots(engine,first["import_id"],tmp_path)
    assert not any(report["counts"].values()) and report["warnings"]
